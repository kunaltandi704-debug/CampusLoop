from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("CampusLoop_Presentation.pptx")
GREEN = RGBColor(35, 90, 60)
INK = RGBColor(30, 48, 38)
MUTED = RGBColor(105, 120, 109)
PAPER = RGBColor(247, 248, 243)
WHITE = RGBColor(255, 255, 255)
LIME = RGBColor(215, 240, 103)
MINT = RGBColor(227, 243, 233)
GOLD = RGBColor(246, 173, 114)


def textbox(slide, x, y, w, h, text, size=18, color=INK, bold=False, font="Aptos", align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(8 if size < 20 else 11)
        if align is not None:
            paragraph.alignment = align
    return shape


def box(slide, x, y, w, h, fill=WHITE, radius=True, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
    return shape


def base(title, kicker, number, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = GREEN if dark else PAPER
    title_color = WHITE if dark else INK
    textbox(slide, 0.72, 0.38, 11.8, 0.3, kicker.upper(), 10, LIME if dark else GREEN, True)
    textbox(slide, 0.72, 0.75, 11.8, 0.62, title, 29, title_color, True, "Aptos Display")
    textbox(slide, 0.72, 7.12, 8, 0.2, "CAMPUSLOOP  ·  HACKATHON DEMO", 8, RGBColor(201, 216, 203) if dark else MUTED, True)
    textbox(slide, 12.0, 7.10, 0.5, 0.2, f"{number:02d}", 9, RGBColor(201, 216, 203) if dark else MUTED, True, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, body, accent=GREEN, dark=False):
    box(slide, x, y, w, h, RGBColor(44, 104, 70) if dark else WHITE)
    textbox(slide, x + 0.23, y + 0.2, w - 0.46, 0.36, heading, 16, LIME if dark else accent, True)
    textbox(slide, x + 0.23, y + 0.7, w - 0.46, h - 0.85, body, 12, RGBColor(231, 239, 232) if dark else MUTED)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# 1 — Cover
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = GREEN
box(slide, 9.1, -0.6, 4.8, 4.8, RGBColor(43, 103, 68))
box(slide, 10.1, 3.3, 4.5, 4.5, RGBColor(30, 77, 52))
textbox(slide, 0.9, 0.75, 6.8, 0.35, "CAMPUS ENGAGEMENT · CREATIVE REWARDS · GEMMA 4", 10, LIME, True)
textbox(slide, 0.9, 1.65, 7.7, 1.3, "CampusLoop", 54, WHITE, True, "Aptos Display")
textbox(slide, 0.94, 3.0, 7.4, 1.05, "Make campus count.\nEarn. Create. Belong.", 25, WHITE, True)
textbox(slide, 0.94, 5.6, 6.8, 0.65, "A rewards loop connecting students, campus moments, creative contributions, and the canteen.", 16, RGBColor(216, 231, 218))
box(slide, 9.15, 1.45, 3.1, 3.45, LIME)
textbox(slide, 9.45, 1.85, 2.5, 0.4, "YOUR CAMPUS LOOP", 11, GREEN, True, align=PP_ALIGN.CENTER)
textbox(slide, 9.55, 2.55, 2.3, 1.8, "SHOW UP\nCREATE\nREDEEM", 22, GREEN, True, "Aptos Display", PP_ALIGN.CENTER)
textbox(slide, 0.94, 7.05, 8, 0.2, "HACKATHON CONCEPT · DEMO ACTIVITIES ARE SAMPLE DATA", 8, RGBColor(201, 216, 203), True)

# 2 — Problem / solution
slide = base("Campus energy is everywhere. Recognition isn't.", "01 · The opportunity", 2)
card(slide, 0.75, 1.7, 3.8, 3.85, "TODAY'S GAP", "Events, volunteering, student design, and canteen rewards are fragmented.\n\nParticipation can go unnoticed. Small creative contributions lack a clear path to recognition.")
card(slide, 4.78, 1.7, 3.8, 3.85, "CAMPUSLOOP", "One student-first place to discover activities, submit original artwork, earn transparent tokens, and redeem canteen treats.")
card(slide, 8.8, 1.7, 3.8, 3.85, "THE RESULT", "More visible opportunities.\n\nClear review and reward rules.\n\nA practical engagement loop that students can understand.", GOLD)
textbox(slide, 1.05, 6.05, 11.2, 0.42, "A demo platform to make participation easier to discover and contributions easier to value.", 16, GREEN, True, align=PP_ALIGN.CENTER)

# 3 — Journey
slide = base("A simple loop from discovery to delight", "02 · Student journey", 3)
steps = [
    ("01", "Discover", "Find workshops, fests, competitions, and volunteering."),
    ("02", "Participate", "Request review; approved participation earns the configured activity reward."),
    ("03", "Create", "Submit an event poster, banner, or canteen promotion."),
    ("04", "Earn", "See every approved credit and spend in a persistent wallet ledger."),
    ("05", "Redeem", "Use tokens for a canteen refreshment or place an order."),
]
for index, (number, heading, body) in enumerate(steps):
    x = 0.75 + index * 2.52
    box(slide, x, 2.15, 2.23, 2.9, WHITE)
    textbox(slide, x + 0.2, 2.38, 0.55, 0.4, number, 12, GREEN, True)
    textbox(slide, x + 0.2, 2.95, 1.85, 0.4, heading, 18, INK, True)
    textbox(slide, x + 0.2, 3.55, 1.85, 1.2, body, 11, MUTED)
    if index < 4:
        textbox(slide, x + 2.23, 3.15, 0.27, 0.35, "→", 18, GREEN, True, align=PP_ALIGN.CENTER)
textbox(slide, 1.0, 5.75, 11.2, 0.4, "Every token change is validated on the server and recorded with its reason.", 15, GREEN, True, align=PP_ALIGN.CENTER)

# 4 — Wallet
slide = base("A wallet students can understand", "03 · Tokens & canteen", 4)
box(slide, 0.8, 1.7, 4.5, 4.6, GREEN)
textbox(slide, 1.15, 2.05, 3.8, 0.35, "BASE REWARD CONVERSION", 10, LIME, True)
textbox(slide, 1.15, 2.75, 3.8, 0.75, "100 tokens = ₹10", 30, WHITE, True, "Aptos Display")
textbox(slide, 1.15, 3.7, 3.7, 1.0, "Demo smoothie\n100 tokens · ₹10 reward value", 17, WHITE, True)
textbox(slide, 1.15, 5.25, 3.7, 0.7, "Other reward prices are set by the admin.", 12, RGBColor(218, 231, 220))
card(slide, 5.7, 1.7, 3.05, 2.0, "EARN", "Approved participation\nApproved creative work", GREEN)
card(slide, 9.0, 1.7, 3.05, 2.0, "SPEND", "Reward catalogue\nCanteen orders", GOLD)
card(slide, 5.7, 3.95, 6.35, 2.35, "TRUST THE LEDGER", "Atomic debit on valid checkout · no negative balance · idempotent requests · eligible pending-order refund exactly once.\n\nStudent history is read from SQLite, never invented by AI.")

# 5 — Contributions
slide = base("Student creativity becomes a recognized contribution", "04 · Poster & banner rewards", 5)
card(slide, 0.8, 1.75, 3.75, 3.75, "SUBMIT", "Original event poster, banner, or canteen promotion.\n\nPNG, JPG, WebP · 5 MB maximum · private owner/admin image access.")
card(slide, 4.78, 1.75, 3.75, 3.75, "REVIEW", "Admin checks the submission and approves or rejects it.\n\nEvery contribution has an explicit Pending Review / Approved / Rejected status.")
card(slide, 8.76, 1.75, 3.75, 3.75, "REWARD ONCE", "Admin configures the token award.\n\nApproval and ledger credit share one database transaction; repeat awards are blocked.")
textbox(slide, 1.1, 6.0, 11.1, 0.4, "Students can always see their own submission history, review status, and earned tokens.", 15, GREEN, True, align=PP_ALIGN.CENTER)

# 6 — Calendar
slide = base("Turn campus dates into visible opportunities", "05 · Academic calendar", 6)
box(slide, 0.85, 1.65, 7.1, 4.85, WHITE)
textbox(slide, 1.15, 1.95, 6.5, 0.3, "ACTIVITY EXPLORER", 10, GREEN, True)
calendar_items = [
    ("UPCOMING", "Sustainable Campus Workshop", "Workshop · hands-on learning"),
    ("UPCOMING", "Inter-Department Design Sprint", "Competition · team challenge"),
    ("UPCOMING", "Campus Fest Poster Jam", "College Fest · creative opportunity"),
]
for i, (tag, title, detail) in enumerate(calendar_items):
    y = 2.55 + i * 1.05
    box(slide, 1.15, y, 0.85, 0.55, MINT)
    textbox(slide, 1.2, y + 0.16, 0.75, 0.2, f"DAY {i+3}", 8, GREEN, True, align=PP_ALIGN.CENTER)
    textbox(slide, 2.25, y, 4.9, 0.3, title, 14, INK, True)
    textbox(slide, 2.25, y + 0.38, 4.9, 0.25, f"{tag} · {detail}", 10, MUTED)
card(slide, 8.35, 1.7, 3.95, 2.0, "ADMIN CONTROLS", "Create, edit, delete, and configure the optional activity reward.")
card(slide, 8.35, 3.95, 3.95, 2.0, "STUDENT VIEW", "Search, category filters, list and month calendar. Dates drive Upcoming / Ongoing / Completed.", GOLD)
textbox(slide, 1.05, 6.72, 11.2, 0.2, "Seeded activity examples are explicitly labeled demo samples—not verified college events.", 10, MUTED, True, align=PP_ALIGN.CENTER)

# 7 — AI architecture
slide = base("Gemma 4 helps explain the loop—not control it", "06 · AI architecture & guardrails", 7, True)
card(slide, 0.85, 1.75, 3.55, 2.5, "STUDENT QUESTION", "Ask about rules, affordable rewards, activities, creative opportunities, or own order status.", dark=True)
card(slide, 4.88, 1.75, 3.55, 2.5, "VERIFIED CONTEXT", "Backend selects only the relevant signed-in student's balance, prices, activity rows, or own orders.", dark=True)
card(slide, 8.9, 1.75, 3.55, 2.5, "GEMMA 4", "Exact model: gemma-4-26b-a4b-it\nGoogle GenAI API · server-side key only", dark=True)
textbox(slide, 4.28, 2.8, 0.45, 0.4, "→", 20, LIME, True, align=PP_ALIGN.CENTER)
textbox(slide, 8.3, 2.8, 0.45, 0.4, "→", 20, LIME, True, align=PP_ALIGN.CENTER)
box(slide, 1.3, 4.8, 10.7, 1.15, RGBColor(44, 104, 70))
textbox(slide, 1.6, 5.1, 10.1, 0.55, "No model fallback · no invented account facts · no direct token/order/submission actions · explicit unavailable error", 15, WHITE, True, align=PP_ALIGN.CENTER)
textbox(slide, 1.1, 6.35, 11.1, 0.35, "AI is a read-only guide. SQLite records and validated backend routes remain authoritative.", 12, RGBColor(216, 231, 218), align=PP_ALIGN.CENTER)

# 8 — Orders
slide = base("Canteen orders are persisted, priced, and refundable by rule", "07 · Order operations", 8)
steps = [
    ("Cart", "Menu item IDs + quantities"),
    ("Validate", "Server price + balance"),
    ("Commit", "Order, items, ledger debit"),
    ("Fulfil", "Pending → Preparing → Ready → Completed"),
]
for i, (heading, body) in enumerate(steps):
    x = 0.9 + i * 3.1
    card(slide, x, 2.1, 2.65, 2.4, heading, body, GOLD if i == 3 else GREEN)
    if i < 3:
        textbox(slide, x + 2.68, 3.0, 0.36, 0.45, "→", 21, GREEN, True, align=PP_ALIGN.CENTER)
box(slide, 1.05, 5.1, 11.2, 0.95, MINT)
textbox(slide, 1.35, 5.35, 10.6, 0.38, "Cancellation: only while Pending. Status and refund are atomic; a unique ledger guard prevents double refunds.", 14, GREEN, True, align=PP_ALIGN.CENTER)
textbox(slide, 1.3, 6.38, 10.6, 0.3, "Students see only their own database-backed order history.", 12, MUTED, align=PP_ALIGN.CENTER)

# 9 — Demo and impact
slide = base("A demo that shows the full loop", "08 · Demo plan & intended impact", 9)
card(slide, 0.8, 1.7, 5.7, 3.9, "LIVE WALKTHROUGH", "1 · Student wallet + 100-token smoothie\n2 · Place order; follow status in Admin\n3 · Browse sample activities; request participation\n4 · Submit artwork; admin configures and approves reward\n5 · Ask Gemma about affordable rewards / order status", GREEN)
card(slide, 6.85, 1.7, 5.7, 3.9, "WHAT TO MEASURE NEXT", "Pilot with one campus canteen and student club.\n\nTrack activity discovery → participation, creative submissions approved, token redemption, repeat visits, and student satisfaction.\n\nTargets belong to a real pilot—not fabricated demo outcomes.", GOLD)
textbox(slide, 1.1, 6.2, 11.2, 0.35, "Seeded balances, accounts, and sample activities are demo data; no real institution is implied.", 12, MUTED, True, align=PP_ALIGN.CENTER)

# 10 — Future
slide = base("A focused foundation—with room to grow", "09 · Impact & future scope", 10, True)
card(slide, 0.85, 1.75, 3.55, 3.7, "NOW", "Transparent wallet & ledger\nCanteen orders & safe refunds\nReviewed creative contributions\nAdmin academic calendar\nVerified-data Gemma assistant", dark=True)
card(slide, 4.88, 1.75, 3.55, 3.7, "NEXT", "Pilot with student groups\nInstitution system integrations\nAccessibility and localization\nOperations analytics\nInstitution identity integration", dark=True)
card(slide, 8.9, 1.75, 3.55, 3.7, "VISION", "A campus where useful participation is discoverable, creativity is valued, and rewards reinforce community—without compromising trust.", dark=True)
textbox(slide, 1.05, 6.15, 11.2, 0.45, "CampusLoop · Make campus count.", 20, LIME, True, align=PP_ALIGN.CENTER)

# 11 — Attendance
slide = base("Attendance rewards are calculated from finalized class records", "10 · Weekly attendance", 11)
card(slide, 0.8, 1.7, 3.75, 3.85, "RECORD", "Admins configure classes, sections, subjects, sessions, enrollments, and teacher assignments.\n\nAssigned teachers mark every enrolled student Present or Absent; duplicates are blocked.")
card(slide, 4.78, 1.7, 3.75, 3.85, "CALCULATE", "Monday 00:00 → next Monday 00:00 in the configured college timezone.\n\nOnly completed weeks with all sessions finalized can qualify.", GOLD)
card(slide, 8.76, 1.7, 3.75, 3.85, "ISSUE ONCE", "<70%: 0 · 70–<80%: 200 · 80–<90%: 300 · 90–100%: 500 tokens.\n\nHighest tier only. Unique student/week reward and ledger entry commit atomically.")
textbox(slide, 1.05, 6.05, 11.2, 0.45, "Corrections are reasoned and audited. Post-award changes require an explicit admin adjustment; historical ledger entries are never rewritten.", 13, GREEN, True, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Created {OUT}")
